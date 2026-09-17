from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN

OUT=Path(__file__).parent
prs=Presentation(r'C:\Users\ASUS\Downloads\SIH2026-IDEA-Presentation-Format (1).pptx')
sid=prs.slides._sldIdLst[-1]; prs.part.drop_rel(sid.rId); prs.slides._sldIdLst.remove(sid)
NAVY='143454'; BLUE='006EB9'; TEAL='008C91'; INK='26384B'; GRAY='617284'; PALE='EDF5FC'; GREEN='EAF6F1'; GOLD='FFF3DA'; WHITE='FFFFFF'
def box(sl,x,y,w,h,fill=PALE,line=None,rounded=True):
 s=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
 s.fill.solid(); s.fill.fore_color.rgb=RGBColor.from_string(fill)
 s.line.fill.background() if not line else None
 if line: s.line.color.rgb=RGBColor.from_string(line)
 if rounded: s.adjustments[0]=0.12
 return s
def text(sl,x,y,w,h,txt,size=17,color=INK,bold=False,align=None):
 s=sl.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf=s.text_frame; tf.word_wrap=True
 tf.margin_left=tf.margin_right=Inches(.025); tf.margin_top=tf.margin_bottom=0
 for i,line in enumerate(txt.split('\n')):
  p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=line; p.font.name='Aptos'; p.font.size=Pt(size); p.font.bold=bold; p.font.color.rgb=RGBColor.from_string(color); p.space_after=Pt(7)
  if align is not None: p.alignment=align
 return s
def card(sl,x,y,w,h,title,body,fill=PALE):
 box(sl,x,y,w,h,fill); text(sl,x+.18,y+.14,w-.36,.38,title,19,NAVY,True)
 offset=.55 if h<=1.5 else .65
 text(sl,x+.18,y+offset,w-.36,h-offset-.08,body,15 if h<=1.5 else 16)
def arrow(sl,x,y,w=.3,h=.25):
 s=sl.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h)); s.fill.solid(); s.fill.fore_color.rgb=RGBColor.from_string(TEAL); s.line.fill.background()
def section(sl,x,y,w,title): text(sl,x,y,w,.4,title,20,BLUE,True)
def strip(sl,y,txt): box(sl,.45,y,12.4,.48,NAVY,rounded=False); text(sl,.62,y+.08,12.05,.35,txt,15,WHITE)
def link(sl,x,y,w,label,url):
 s=text(sl,x,y,w,.32,label,13,BLUE); s.text_frame.paragraphs[0].runs[0].hyperlink.address=url

# Retain the supplied logo, footer and team badge; replace only content areas.
for i,sl in enumerate(prs.slides):
 if i==0:
  for s in list(sl.shapes):
   if s.shape_type==13 and s.left>Inches(10): continue
   s._element.getparent().remove(s._element)
 else:
  for j in [2,1]:
   s=sl.shapes[j]; s._element.getparent().remove(s._element)
  for s in sl.shapes:
   if s.has_text_frame and 'Your Team Name' in s.text:
    s.text='TEAM\n[TBD]'
    for p in s.text_frame.paragraphs:
     p.font.name='Aptos'; p.font.size=Pt(15); p.font.bold=True; p.alignment=PP_ALIGN.CENTER
 titles=['','IDEA TITLE','TECHNICAL APPROACH','FEASIBILITY AND VIABILITY','IMPACT AND BENEFITS','RESEARCH AND REFERENCES']
 if i: text(sl,1.9,.48,8.6,.65,titles[i],28,NAVY,True,PP_ALIGN.CENTER)

s=prs.slides[0]
text(s,.5,.3,9.8,.6,'SMART INDIA HACKATHON 2026',26,BLUE,True)
text(s,.55,1.35,8,.8,'NETRA-MAIL',48,NAVY,True)
text(s,.58,2.25,7.3,.85,'From a suspicious email to an\ninvestigation-ready forensic report.',25,INK)
box(s,.55,3.35,7.2,2.95,PALE)
text(s,.78,3.55,6.75,2.65,'Problem Statement ID: 26106 — confirm\nPS Title: [Exact portal title to be confirmed]\nTheme: [Confirm portal theme]\nPS Category: Software\nTeam ID: [To be added]\nTeam Name: [Registered name to be added]',17)
for j,(a,b) in enumerate([('DETECT','Explainable phishing risk'),('INVESTIGATE','Message, URL and relay evidence'),('PRESERVE','Case timeline and forensic report')]):
 card(s,8.15,1.6+j*1.55,4.6,1.3,a,b,[PALE,GREEN,GOLD][j])
strip(s,6.65,'User-initiated Gmail analysis  •  Explainable findings  •  Tamper-evident evidence')
s.notes_slide.notes_text_frame.text='Confirm the official problem-statement title, theme, team ID and registered team name before submission. PS 26106 comes from the project documentation and requires portal confirmation.'

s=prs.slides[1]
section(s,.45,1.4,12,'Proposed Solution (Describe your Idea/Solution/Prototype)')
text(s,.47,1.9,12.2,.5,'NETRA-Mail connects a Gmail risk badge to a complete SOC investigation.',22,NAVY,True)
section(s,.45,2.65,3.8,'Detailed explanation')
section(s,4.7,2.65,3.9,'How it addresses the problem')
section(s,8.95,2.65,3.9,'Innovation and uniqueness')
card(s,.45,3.15,3.93,2.55,'One connected workflow','• Analyze the selected email.\n• Explain sender, text and link risks.\n• Open the matching forensic report.')
card(s,4.7,3.15,3.93,2.55,'Evidence behind the alert','• Makes warning signals readable.\n• Reduces fragmented investigation.\n• Preserves context for follow-up.',GREEN)
card(s,8.95,3.15,3.93,2.55,'Evidence-aware decisions','• Labels missing authentication data.\n• Connects related indicators.\n• Combines risk, trace and custody.',GOLD)
strip(s,6.15,'Analyze current email  →  Risk + reasons  →  View forensic report  →  Analyst action')

s=prs.slides[2]
section(s,.45,1.4,12,'Methodology and process for implementation')
nodes=[('01  INPUT','Manual Gmail scan\nor original EML'),('02  ANALYZE','Text + ML + URLs\nheaders + attachments'),('03  EXPLAIN','Risk + findings\nconfidence + limits'),('04  TRACE','Relay / IP context\nrelated indicators'),('05  PRESERVE','Case + SHA-256\nreport + custody')]
for j,(title,body) in enumerate(nodes):
 x=.45+j*2.52; card(s,x,2.02,2.28,1.65,title,body,PALE if j%2==0 else GREEN)
 if j<4: arrow(s,x+2.3,2.7,.2,.22)
box(s,.45,3.94,12.4,.65,GOLD); text(s,.63,4.1,12,.4,'Evidence gate: original bytes enable signature checks; reconstructed Gmail content is labelled incomplete.',16,NAVY,True)
section(s,.45,4.9,12,'Technologies to be used')
for j,(a,b) in enumerate([('INTERFACE','JavaScript MV3\nStreamlit dashboard'),('API + INTELLIGENCE','Python · FastAPI\nscikit-learn · DNS'),('EVIDENCE','SQLite · SHA-256\nAES-GCM encryption'),('DEPLOYMENT','Render API + dashboard\nHTTPS · scoped API keys')]):
 card(s,.45+j*3.15,5.4,2.95,1.24,a,b,PALE)

s=prs.slides[3]
section(s,.45,1.4,5.6,'Analysis of the feasibility of the idea')
card(s,.45,1.97,5.55,1.42,'Technical feasibility','Working Gmail extension, FastAPI analysis,\nStreamlit investigations and report generation.')
card(s,.45,3.6,5.55,1.42,'Operational and financial viability','Pilot within one institution; open-source stack.\nHosting, storage and intelligence have running costs.',GREEN)
box(s,.45,5.25,5.55,1.35,NAVY); text(s,.68,5.42,5.1,.42,'149 automated tests passed',27,WHITE,True); text(s,.68,6.03,5.1,.35,'Local regression evidence; not detection accuracy.',14,WHITE)
section(s,6.4,1.4,6.5,'Potential challenges and risks → Strategies')
for j,(title,body) in enumerate([('Incomplete message headers','Fetch original MIME / upload EML; label unavailable checks.'),('False positives and model drift','Analyst review; evaluate on unseen data before rollout.'),('Privacy and online access','Scoped identities, encryption, retention and dashboard login.'),('Provider outages and persistence','Show unavailable signals; durable storage and backups.')]):
 y=1.97+j*1.18; box(s,6.4,y,6.43,1.02,[PALE,GREEN,GOLD,PALE][j]); text(s,6.6,y+.1,6,.31,title,18,NAVY,True); text(s,6.6,y+.49,6,.48,body,15)
s.notes_slide.notes_text_frame.text='149 automated tests passed in the local project run recorded on 15 September 2026. This is a software regression count, not classification accuracy. Before a real institutional deployment, implement dashboard authentication, verify backups and test encrypted evidence retention. Provider-backed checks depend on configured external services.'

s=prs.slides[4]
section(s,.45,1.4,12,'Potential impact on the target audience')
for j,(a,b) in enumerate([('EMAIL USERS','Understand why a message is risky\nand open its supporting report.'),('SOC / IT TEAMS','Triage findings in one place;\ncorrelate recurring indicators.'),('INSTITUTIONS','Build consistent incident records\nfor review and response.')]): card(s,.45+j*4.2,2,3.95,1.65,a,b,[PALE,GREEN,GOLD][j])
section(s,.45,3.95,12,'Benefits of the solution (social, economic, environmental, etc.)')
for j,(a,b) in enumerate([('SOCIAL','Clear warnings support phishing\nawareness and informed decisions.'),('ECONOMIC','Aim to reduce review effort and\nfraud exposure; validate in pilots.'),('RESOURCE USE','Digital reports reduce paper use;\nmanual scans limit unnecessary work.')]): card(s,.45+j*4.2,4.5,3.95,1.4,a,b,[PALE,GREEN,GOLD][j])
strip(s,6.18,'Pilot success measures: false-positive rate  •  phishing recall  •  review time  •  report completeness')

s=prs.slides[5]
section(s,.45,1.4,12,'Details / Links of the reference and research work')
refs=[('01  Email authentication','RFC 6376 — DKIM signatures','Original-byte integrity and signing-domain checks.','https://www.rfc-editor.org/info/rfc6376/'),('02  Sender authorization','RFC 7208 — Sender Policy Framework','SPF needs trusted SMTP connection context.','https://www.rfc-editor.org/info/rfc7208/'),('03  Domain alignment','RFC 7489 — DMARC','Connect authentication evidence to the From domain.','https://www.rfc-editor.org/info/rfc7489/'),('04  Forensic methodology','NIST SP 800-86','Collection, examination, analysis and reporting.','https://csrc.nist.gov/pubs/sp/800/86/final')]
for j,(a,b,c,u) in enumerate(refs):
 x=.45+(j%2)*6.35; y=2+(j//2)*1.68; box(s,x,y,6.05,1.47,PALE if j%2==0 else GREEN); text(s,x+.17,y+.12,5.7,.32,a,18,NAVY,True); text(s,x+.17,y+.53,5.7,.32,c,15); link(s,x+.17,y+1.01,5.7,b,u)
box(s,.45,5.58,12.4,1.03,GOLD); text(s,.65,5.7,11.9,.33,'Prototype and validation evidence',19,NAVY,True)
link(s,.65,6.12,11.9,'NETRA-MAIL source repository • implementation, tests and deployment setup','https://github.com/nitheenkumar25bcy37-alt/NETRA-MAIL')
s.notes_slide.notes_text_frame.text='Reference links are clickable. RFC 7489 is cited as the DMARC basis used in this implementation; consult RFC Editor for subsequent updates. Layout inspiration: user-supplied 2025 Sarthi / Bit-Storm reference PDF. Content and diagrams describe NETRA-Mail. No independent phishing-accuracy or financial-impact claim is made.'
if __name__ == '__main__':
 prs.save(OUT/'NETRA_MAIL_SIH_2026.pptx')
 print('Saved six-slide presentation')

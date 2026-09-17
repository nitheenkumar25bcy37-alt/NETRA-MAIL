import runpy
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

OUT=Path(__file__).parent
d=runpy.run_path(str(OUT/'build_deck.py'))
prs=d['prs']; box=d['box']; text=d['text']; arrow=d['arrow']; link=d['link']
NAVY='173958'; BLUE='006EB9'; TEAL='008B8B'; INK='233746'; WHITE='FFFFFF'
colors=['E1EFFC','DDF5F1','EAF3D7','EEE4F8','FFF1C9','FFE5D2']
def t(s,x,y,w,h,content,size=13,bold=False,color=INK,center=False):
 sh=text(s,x,y,w,h,content,size,color,bold,PP_ALIGN.CENTER if center else None)
 for p in sh.text_frame.paragraphs: p.space_after=Pt(3); p.space_before=Pt(0)
 return sh
def heading(s,x,y,w,label):
 t(s,x,y,w,.3,label,18,True,NAVY)
 box(s,x,y+.34,w,.018,BLUE,rounded=False)
def panel(s,x,y,w,h,title,body,col=0,size=13):
 box(s,x,y,w,h,colors[col],rounded=True)
 t(s,x+.13,y+.1,w-.26,.34,title,15,True,NAVY)
 t(s,x+.13,y+.43,w-.26,h-.5,body,size)
def conn(s,x1,y1,x2,y2):
 sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2)); sh.line.color.rgb=RGBColor.from_string('4C7696'); sh.line.width=Pt(1.4)
def bullet(s,x,y,w,label,body,h=.52):
 t(s,x,y,w,.24,label,13,True,NAVY); t(s,x,y+.25,w,h-.2,body,12.6)
def footer_note(s,txt): t(s,.43,6.62,12.45,.23,txt,10.5,False,'607184')

# Rebuild body content while retaining the official 2026 branding and footer.
titles=['','NETRA-MAIL | PROPOSED SOLUTION','TECHNICAL APPROACH','FEASIBILITY AND VIABILITY','IMPACT AND BENEFITS','RESEARCH AND REFERENCES']
for i in range(1,6):
 s=prs.slides[i]
 for sh in list(s.shapes):
  keep=(sh.shape_type==13 or (sh.has_text_frame and (sh.text=='TEAM\n[TBD]' or sh.text.strip()==str(i+1) or '@SIH' in sh.text)) or (sh.top>=Inches(6.9)))
  if not keep: sh._element.getparent().remove(sh._element)
 t(s,1.8,.34,8.85,.67,titles[i],27,True,NAVY,True)

# Title: emphasize distinctive capabilities without inventing registration details.
s=prs.slides[0]
for sh in s.shapes:
 if sh.has_text_frame:
  if 'From a suspicious email' in sh.text:
   sh.text='Multilingual phishing intelligence.\nInvestigation-ready email forensics.'
   for p in sh.text_frame.paragraphs: p.font.name='Aptos'; p.font.size=Pt(25); p.font.color.rgb=RGBColor.from_string(INK)
  if sh.text=='Explainable phishing risk':
   sh.text='Multilingual text, URLs, QR and attachments'
   for p in sh.text_frame.paragraphs: p.font.name='Aptos'; p.font.size=Pt(16); p.font.color.rgb=RGBColor.from_string(INK)
  if 'User-initiated Gmail analysis  •' in sh.text:
   sh.text='6 language profiles  •  Explainable risk  •  Origin context  •  Tamper-evident forensic reporting'
   for p in sh.text_frame.paragraphs: p.font.name='Aptos'; p.font.size=Pt(15); p.font.color.rgb=RGBColor.from_string(WHITE)

s=prs.slides[1]
heading(s,.35,1.2,6.05,'Challenges and problems')
for j,(a,b) in enumerate([
 ('Language barriers in phishing detection','Regional scripts, transliteration and mixed-language messages need contextual analysis.'),
 ('Deception across multiple email surfaces','Spoofed senders, credential links, payment requests and image-based lures obscure intent.'),
 ('Fragmented investigation and missing context','A risk label alone does not explain relay paths, related indicators or evidence history.')]):
 bullet(s,.44,1.68+j*.65,5.83,a,b,.61)
heading(s,.35,3.72,6.05,'Proposed solution — detailed explanation')
for j,(a,b) in enumerate([
 ('1. Multilingual phishing intelligence','Combine language/script cues with urgency, OTP, credential and financial-fraud indicators.'),
 ('2. Multi-layer email inspection','Fuse text/ML findings, sender alignment, domain/URL checks, static attachments and QR/OCR signals.'),
 ('3. Investigation and evidence lifecycle','Open the exact email in the SOC dashboard; inspect findings, correlate cases and export reports.')]):
 bullet(s,.44,4.19+j*.71,5.82,a,b,.69)
heading(s,6.68,1.2,6.27,'Innovation and uniqueness')
for j,(a,b) in enumerate([
 ('MULTILINGUAL','English + 5 Indian\nlanguage profiles'),('EXPLAINABLE','Risk score + reasons\n+ evidence limits'),('FORENSIC','SHA-256 + custody\n+ encrypted evidence'),('CONNECTED','Gmail badge → exact\nSOC investigation')]):
 x=6.7+(j%2)*3.18; y=1.71+(j//2)*1.1
 panel(s,x,y,3.02,.96,a,b,j,12.4)
heading(s,6.68,4.03,6.27,'How it addresses the problem')
bullet(s,6.79,4.49,5.99,'Language-aware review','Native scripts + Romanized cues; language is context, never a malicious verdict.',.57)
bullet(s,6.79,5.14,5.99,'From alert to evidence','Relay/IP context and related indicators help analysts investigate recurring campaigns.',.57)
bullet(s,6.79,5.79,5.99,'Controlled and transparent analysis','Manual scan, scoped access and explicit labels when original headers are unavailable.',.57)
footer_note(s,'Language profiles: English · Hindi · Tamil · Telugu · Kannada · Malayalam. Per-language accuracy remains to be measured.')

s=prs.slides[2]
heading(s,.35,1.2,6.14,'Methodology and process')
methods=[('01  Acquire & authenticate','User selects Gmail content or submits EML; API checks identity, origin and request limits.'),('02  Normalize & extract','Parse MIME, sender fields, relay headers, URLs and static attachments; mask sensitive fields.'),('03  Analyze multilingual signals','Detect script / Romanized cues; inspect urgency, credentials, financial fraud and mixed scripts.'),('04  Fuse evidence & explain','Combine ML, rules, URL/domain checks and authentication evidence with confidence and limitations.'),('05  Investigate & preserve','Link indicators and cases, store encrypted evidence, record custody and generate forensic reports.')]
for j,(a,b) in enumerate(methods):
 y=1.68+j*.57
 t(s,.46,y,5.93,.23,a,12.5,True,NAVY)
 t(s,.46,y+.23,5.93,.34,b,11.2)
heading(s,.35,4.61,6.14,'Technology stack — mapped to the prototype')
stack=[('Py','Python','Analysis engine','3776AB'),('API','FastAPI','REST + validation','009688'),('JS','JavaScript','Chrome MV3','A48000'),('ML','scikit-learn','Risk models','F28B20'),('ST','Streamlit','SOC interface','D94D60'),('DB','SQLite','Forensic ledger','236B8E'),('CV','OpenCV','QR decoding','5F8A35'),('OCR','RapidOCR','Image text','8060AF'),('DNS','dkimpy / DNS','Authentication','008B8B'),('AES','cryptography','AES-GCM','315178'),('CTR','Docker','Isolated worker*','1680B9'),('R','Render','Hosted services','343C48')]
for j,(glyph,name,desc,col) in enumerate(stack):
 x=.4+(j%4)*1.53; y=5.1+(j//4)*.47
 icon=OUT/'stack_icons'/f'{j}.png'
 if icon.exists(): s.shapes.add_picture(str(icon),Inches(x),Inches(y),width=Inches(.34),height=Inches(.34))
 else:
  box(s,x,y,.35,.34,col); t(s,x+.015,y+.07,.32,.2,glyph,9,True,WHITE,True)
 t(s,x+.41,y,1.1,.2,name,10.5,True,NAVY); t(s,x+.41,y+.21,1.1,.22,desc,8.7)
heading(s,6.76,1.2,6.19,'System architecture and working flow')
panel(s,6.83,1.72,2.05,1.18,'INPUT LAYER','Gmail extension\nOriginal EML / MIME\nScoped API key',0,12)
panel(s,9.28,1.72,3.57,1.18,'FASTAPI GATEWAY','Origin + identity validation\nSize limits · ingestion · orchestration',5,12)
arrow(s,8.94,2.19,.27,.22)
conn(s,11.07,2.91,11.07,3.15)
panel(s,9.28,3.15,3.57,1.56,'MULTI-LAYER ANALYSIS','Multilingual NLP + ML scoring\nSender / DKIM / DMARC evidence\nURL + domain + IP intelligence\nAttachment / QR / OCR signals',0,12)
panel(s,6.83,3.15,2.05,1.56,'SOURCES','Received headers\nDNS + public IPs\nOriginal bytes\nOptional reputation*',1,11.5)
arrow(s,8.94,3.8,.27,.22)
conn(s,11.07,4.72,11.07,4.97)
panel(s,9.28,4.97,3.57,1.23,'SOC INVESTIGATION','Findings · relay trace · graph\nCases · custody · PDF/HTML reports',3,12)
panel(s,6.83,4.97,2.05,1.23,'PRESERVATION','SQLite + SHA-256\nAES-GCM evidence\nVersioned records',2,11.5)
conn(s,8.88,5.57,9.28,5.57)
footer_note(s,'*Isolated inspection / reputation require separate configuration. Missing original bytes or provider data are labelled, not assumed safe.')

s=prs.slides[3]
heading(s,.35,1.2,7.76,'Analysis of feasibility and viability')
box(s,3.06,1.73,2.28,.61,colors[0],line='4C7696'); t(s,3.15,1.84,2.1,.4,'NETRA-MAIL PILOT',16,True,NAVY,True)
conn(s,4.2,2.34,4.2,2.55); conn(s,1.29,2.55,7.16,2.55)
branches=[('TECHNICAL',0,['Existing Python / ML stack; modular analyzers and REST APIs.','Six language profiles; Romanized cues and explainable findings.','149 local regression tests passed; unseen-data validation next.']),('FINANCIAL',1,['Open-source components reduce initial licensing overhead.','Budget for compute, durable evidence storage and provider usage.','Institutional pilot first; scale only after measured value.']),('OPERATIONAL',2,['Manual Gmail scan and direct report link fit analyst workflows.','Scoped roles, case notes and custody support team review.','Train analysts; define retention, backups and response procedures.']),('SOCIAL',3,['Regional-language cues support more inclusive phishing review.','Readable reasons build awareness of social-engineering tactics.','Human review and explicit limits reduce false confidence.'])]
for j,(a,c,items) in enumerate(branches):
 x=.4+j*1.96; conn(s,x+.87,2.55,x+.87,2.76); box(s,x,2.76,1.78,.49,colors[c],line='7392A9'); t(s,x+.06,2.89,1.66,.23,a,13,True,NAVY,True)
 for k,body in enumerate(items):
  box(s,x,3.43+k*.93,1.78,.85,colors[c]); t(s,x+.09,3.52+k*.93,1.6,.7,body,10.5)
heading(s,8.44,1.2,4.52,'Risks and mitigation strategies')
panel(s,8.44,1.75,4.52,1.45,'TECHNICAL CHALLENGES','• Missing originals → EML / authorized MIME fetch.\n• Language/model errors → labelled per-language tests.\n• Provider outages → timeouts and unavailable labels.',2,12)
panel(s,8.44,3.36,4.52,1.34,'FINANCIAL CHALLENGES','• Hosting/storage costs → staged institutional pilot.\n• Paid intelligence → optional, budgeted providers.\n• Data growth → retention and backup policy.',4,12)
panel(s,8.44,4.86,4.52,1.4,'OPERATIONAL CHALLENGES','• Public access → dashboard login before real use.\n• False alerts → analyst review and feedback.\n• Evidence continuity → durable disk + restore drills.',5,12)
footer_note(s,'Pilot readiness is distinct from production readiness. Regression tests do not establish phishing accuracy or legal admissibility.')

s=prs.slides[4]
heading(s,.35,1.2,6.1,'Potential impact on the target audience')
impacts=[('Citizens and institutional email users','Language-aware warnings expose credential theft, payment pressure and deceptive links.'),('SOC analysts and campus IT teams','One investigation view connects findings, observed infrastructure, campaigns and evidence.'),('Incident-response and review teams','Versioned records, custody events and forensic reports improve handover and review.'),('Regional-language communities','English, Hindi, Tamil, Telugu, Kannada and Malayalam cues broaden contextual coverage.')]
for j,(a,b) in enumerate(impacts): bullet(s,.45,1.71+j*.81,5.88,a,b,.76)
heading(s,.35,5.04,6.1,'Benefits — social, economic and resource use')
t(s,.46,5.53,5.88,1.02,'• Social: clearer phishing awareness and regional-language inclusion.\n• Economic: aim to reduce manual review effort and fraud exposure.\n• Resource use: digital reports and user-initiated scans avoid unnecessary processing.',13)
heading(s,6.79,1.2,6.15,'Adoption, sustainability and measurable value')
panel(s,6.79,1.72,6.15,1.04,'TARGET USERS / EARLY ADOPTERS','University IT cells, institutional SOCs and small organizations handling suspicious email reports.',0,13)
panel(s,6.79,2.92,6.15,1.17,'VALUE-ADDED CAPABILITIES','Multilingual cues · QR/OCR inspection · explainable risk\nAuthentication evidence · campaign links · traceable reports',1,13)
panel(s,6.79,4.25,6.15,1.0,'SUSTAINABILITY MODEL — PROPOSED','Institution-managed hosting; optional paid intelligence. Evaluate support subscriptions after pilot validation.',4,13)
t(s,6.85,5.48,6,.3,'PHASED ROLLOUT',14,True,NAVY)
for j,a in enumerate(['Institution\npilot','Measure &\ncalibrate','Multi-team\nadoption']):
 x=6.84+j*2.06; box(s,x,5.93,1.84,.58,colors[j]); t(s,x+.06,6,1.72,.45,a,12,True,NAVY,True)
 if j<2: arrow(s,x+1.88,6.1,.14,.18)
footer_note(s,'Measure: per-language precision/recall · false-positive rate · median/p95 scan latency · review time · report completeness. No outcome claims yet.')

s=prs.slides[5]
heading(s,.35,1.2,6.15,'Reference research and standards')
refs=[('DKIM — RFC 6376','Signature verification over original message bytes; signing-domain evidence.','https://www.rfc-editor.org/info/rfc6376/'),('SPF — RFC 7208','Sender authorization requires trusted SMTP connection context.','https://www.rfc-editor.org/info/rfc7208/'),('DMARC — RFC 7489','From-domain alignment using DKIM and trusted SPF evidence.','https://www.rfc-editor.org/info/rfc7489/'),('Digital forensics — NIST SP 800-86','Collection → examination → analysis → reporting.','https://csrc.nist.gov/pubs/sp/800/86/final')]
for j,(a,b,u) in enumerate(refs):
 y=1.75+j*1.02; box(s,.4,y,6.06,.88,colors[j]); link(s,.55,y+.1,5.73,a,u); t(s,.55,y+.45,5.73,.34,b,12)
heading(s,6.85,1.2,6.1,'Prototype evidence and validation roadmap')
panel(s,6.85,1.75,6.05,1.25,'IMPLEMENTED PROTOTYPE','Chrome MV3 + FastAPI + Streamlit; multilingual rules;\nQR/OCR integration; authentication evidence; cases,\ncorrelation, encrypted storage and report generation.',0,13)
panel(s,6.85,3.15,6.05,1.32,'MULTILINGUAL DESIGN BASIS','Unicode script detection + contextual security cues.\nNative scripts and Romanized/transliterated patterns.\nOutput: language, method, cues, confidence and limits.',1,13)
panel(s,6.85,4.62,6.05,1.56,'EVALUATION TO COMPLETE','• Unseen benign/phishing corpus with language labels.\n• Native + Romanized + mixed-script cases.\n• False positives, recall and latency by language.\n• Original-MIME and missing-evidence scenarios.\n• Hosted access, persistence and recovery checks.',4,12.5)
heading(s,.35,6.02,6.15,'Project source and reproducible evidence')
link(s,.43,6.43,6.05,'github.com/nitheenkumar25bcy37-alt/NETRA-MAIL','https://github.com/nitheenkumar25bcy37-alt/NETRA-MAIL')
t(s,6.89,6.37,5.93,.33,'149 local automated tests passed; no claimed accuracy.',12,True,NAVY)
for i in range(1,6):
 prs.slides[i].notes_slide.notes_text_frame.text += '\nDetailed revision: multilingual analysis is heuristic contextual detection, not translation or a multilingual UI. Native script labels can be ambiguous for shared scripts. QR/OCR performance is separately provider/decoder dependent. External validation and production hardening remain required.'
prs.save(OUT/'NETRA_MAIL_SIH_2026_Detailed.pptx')
print('Detailed six-slide deck saved')
